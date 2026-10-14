"""End-to-end tests for the checklist backend.

Run from the ``backend`` directory:

    python -m unittest discover -s tests -v

Scratch files go to ``backend/out/_tests`` rather than the system temp
directory, which is not writable in every sandbox.
"""

from __future__ import annotations

import json
import shutil
import sys
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from docx import Document  # noqa: E402

from compliance.catalog import QUESTION_IDS, derive_completion  # noqa: E402
from compliance.generator import (  # noqa: E402
    cell_text,
    completion_metadata,
    generate_document,
    iter_question_rows,
    verify_document,
)
from compliance.server import create_server  # noqa: E402
from compliance.service import ComplianceService  # noqa: E402

TEMPLATE = BACKEND_DIR / "templates" / "China-AI-Compliance-Checklist-IT-Solution-v1.3.docx"
TESTS_ROOT = BACKEND_DIR / "out" / "_tests"
QUESTION_COUNT = 26


def completion_entries(path) -> dict[str, str]:
    """Read the Completion Record table out of a generated document."""

    doc = Document(str(path))
    table = next(
        table
        for table in doc.tables
        if len(table.columns) == 2 and cell_text(table.rows[0].cells[0]) == "Field"
    )
    return {cell_text(row.cells[0]): cell_text(row.cells[1]) for row in table.rows[1:]}


class WorkspaceCase(unittest.TestCase):
    """Give each test a clean directory under ``out/_tests``."""

    def setUp(self) -> None:
        self.workdir = TESTS_ROOT / f"{type(self).__name__}_{self._testMethodName}"
        shutil.rmtree(self.workdir, ignore_errors=True)
        self.workdir.mkdir(parents=True, exist_ok=True)

    def tearDown(self) -> None:
        shutil.rmtree(self.workdir, ignore_errors=True)


class GeneratorTests(WorkspaceCase):
    def test_all_questions_default_to_yes(self) -> None:
        output = self.workdir / "checklist.docx"
        summary = generate_document(TEMPLATE, output)

        self.assertEqual(summary["questionCount"], QUESTION_COUNT)
        self.assertEqual(summary["problems"], [])
        self.assertEqual(verify_document(output), [])

        doc = Document(str(output))
        fonts = {
            run.font.name
            for _, row in iter_question_rows(doc)
            for run in row.cells[2].paragraphs[0].runs
            if run.text.strip()
        }
        self.assertTrue(fonts, "the Yes cells should carry a font")
        self.assertNotIn(None, fonts)

    def test_only_the_yes_column_changes(self) -> None:
        output = self.workdir / "checklist.docx"
        generate_document(TEMPLATE, output)

        before = dict(iter_question_rows(Document(str(TEMPLATE))))
        after = dict(iter_question_rows(Document(str(output))))

        changed = []
        for question_id, row in after.items():
            original = before[question_id]
            for column in range(5):
                if cell_text(original.cells[column]) != cell_text(row.cells[column]):
                    changed.append((question_id, column))

        self.assertEqual(len(changed), QUESTION_COUNT)
        self.assertTrue(all(column == 2 for _, column in changed))

    def test_step0_and_completion_metadata(self) -> None:
        output = self.workdir / "checklist.docx"
        generate_document(
            TEMPLATE,
            output,
            step0="AI-2",
            metadata=completion_metadata(
                name="MAZE purchasing order tracking agent",
                completed_by="THG4SZH / 2026-10-14",
                result="Legal as described",
                storage_location="MP-308.24 project folder",
            ),
        )
        doc = Document(str(output))
        step0 = next(p.text for p in doc.paragraphs if p.text.startswith("Step 0"))
        self.assertIn("\u2611 AI-2", step0)
        self.assertIn("\u2610 AI-1", step0)

        completion = next(
            table
            for table in doc.tables
            if len(table.columns) == 2 and cell_text(table.rows[0].cells[0]) == "Field"
        )
        entries = {
            cell_text(row.cells[0]): cell_text(row.cells[1]) for row in completion.rows[1:]
        }
        self.assertEqual(
            entries["IT-Solution / AI system name"], "MAZE purchasing order tracking agent"
        )
        self.assertIn("[x] Legal as described", entries["Result"])
        self.assertIn("[ ] Action required", entries["Result"])
        self.assertEqual(entries["No answers"], "none")

    def test_unknown_question_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            generate_document(TEMPLATE, self.workdir / "x.docx", {"Z9": "Yes"})


class ServiceTests(WorkspaceCase):
    def setUp(self) -> None:
        super().setUp()
        self.service = ComplianceService(
            storage_dir=self.workdir / "out", template_path=TEMPLATE
        )

    def test_acknowledge_generates_document(self) -> None:
        state = self.service.acknowledge("REQ-1", actor="THG4SZH")

        self.assertTrue(state["termsAcknowledged"])
        self.assertEqual(state["acknowledgedBy"], "THG4SZH")
        self.assertEqual(len(state["answers"]), QUESTION_COUNT)
        self.assertEqual(set(state["answers"].values()), {"Yes"})
        self.assertEqual(state["document"]["problems"], [])
        self.assertTrue(self.service.document_path("REQ-1").is_file())
        self.assertEqual(verify_document(self.service.document_path("REQ-1")), [])

    def test_acknowledge_is_idempotent(self) -> None:
        first = self.service.acknowledge("REQ-2", actor="THG4SZH")
        second = self.service.acknowledge("REQ-2", actor="THG4SZH")
        self.assertEqual(first, second)

    def test_revoke_clears_the_flag_and_keeps_the_file(self) -> None:
        self.service.acknowledge("REQ-3", actor="THG4SZH")
        state = self.service.acknowledge("REQ-3", acknowledged=False)
        self.assertFalse(state["termsAcknowledged"])
        self.assertIsNone(state["document"])
        self.assertTrue(self.service.document_path("REQ-3").is_file())

    def test_bad_request_id_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            self.service.acknowledge("../etc/passwd")

    def test_completion_record_is_filled(self) -> None:
        document = self.service.document_path("REQ-4")
        state = self.service.acknowledge(
            "REQ-4", actor="THG4SZH", name="MAZE purchasing order tracking agent"
        )
        entries = completion_entries(document)

        self.assertEqual(
            entries["IT-Solution / AI system name"], "MAZE purchasing order tracking agent"
        )
        self.assertTrue(entries["Completed by / Date"].startswith("THG4SZH / "))
        self.assertEqual(entries["Section(s) skipped"], "none")
        self.assertEqual(entries["No answers"], "none")
        self.assertEqual(
            entries["Not assessable answers (question IDs and missing information)"], "none"
        )
        self.assertIn("[x] Legal as described", entries["Result"])
        self.assertIn("[ ] Action required", entries["Result"])
        self.assertEqual(entries["Storage location"], str(document.parent))
        self.assertEqual(entries["Reviewer sign-off"], "pending")
        self.assertEqual(state["completion"]["Result"], "Legal as described")

    def test_completion_record_defaults_to_the_request_id(self) -> None:
        self.service.acknowledge("REQ-5", actor="THG4SZH")
        entries = completion_entries(self.service.document_path("REQ-5"))
        self.assertEqual(entries["IT-Solution / AI system name"], "REQ-5")
        self.assertTrue(all("[TBD" not in value for value in entries.values()))


class DerivationTests(unittest.TestCase):
    def test_result_follows_the_verdict_table(self) -> None:
        ids = QUESTION_IDS

        clean = {question_id: "Yes" for question_id in ids}
        self.assertEqual(derive_completion(clean, ids)["result"], "Legal as described")

        with_no = dict(clean)
        with_no["A1"] = "No"
        self.assertEqual(derive_completion(with_no, ids)["result"], "Action required")

        with_na = dict(clean)
        for index in (1, 2, 3):
            with_na[f"B{index}"] = "Not assessable"
        derived = derive_completion(with_na, ids)
        self.assertEqual(derived["result"], "Cannot determine")
        self.assertEqual(derived["skipped"], ["B"])
        self.assertEqual(derived["not_assessable"], ["B1", "B2", "B3"])


class ApiTests(WorkspaceCase):
    def setUp(self) -> None:
        super().setUp()
        service = ComplianceService(
            storage_dir=self.workdir / "out", template_path=TEMPLATE
        )
        self.httpd = create_server(service, host="127.0.0.1", port=0)
        self.port = self.httpd.server_address[1]
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self) -> None:
        self.httpd.shutdown()
        self.httpd.server_close()
        self.thread.join(timeout=5)
        super().tearDown()

    def url(self, path: str) -> str:
        return f"http://127.0.0.1:{self.port}{path}"

    def test_checkbox_tick_generates_the_document(self) -> None:
        request = urllib.request.Request(
            self.url("/api/ai-requests/REQ-API/compliance/acknowledge"),
            data=json.dumps(
                {
                    "termsAcknowledged": True,
                    "actor": "THG4SZH",
                    "name": "REQ-API agent",
                }
            ).encode(),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(request, timeout=30) as response:
            self.assertEqual(response.status, 200)
            payload = json.load(response)

        self.assertTrue(payload["termsAcknowledged"])
        self.assertEqual(payload["answers"]["A1"], "Yes")
        self.assertEqual(payload["document"]["problems"], [])
        self.assertEqual(
            payload["documentUrl"], "/api/ai-requests/REQ-API/compliance/checklist.docx"
        )

        with urllib.request.urlopen(self.url(payload["documentUrl"]), timeout=30) as response:
            body = response.read()
        self.assertTrue(body.startswith(b"PK"), "expected a DOCX (zip) payload")

        saved = self.workdir / "downloaded.docx"
        saved.write_bytes(body)
        self.assertEqual(verify_document(saved), [])

        entries = completion_entries(saved)
        self.assertEqual(entries["IT-Solution / AI system name"], "REQ-API agent")
        self.assertIn("[x] Legal as described", entries["Result"])

    def test_state_endpoint_reports_not_acknowledged(self) -> None:
        with urllib.request.urlopen(
            self.url("/api/ai-requests/REQ-API/compliance"), timeout=30
        ) as response:
            self.assertEqual(response.status, 200)
            payload = json.load(response)
        self.assertFalse(payload["termsAcknowledged"])
        self.assertIsNone(payload["document"])
        self.assertEqual(payload["answers"], {})

    def test_unknown_route_returns_404(self) -> None:
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            urllib.request.urlopen(self.url("/api/ai-requests/REQ-API/other"), timeout=30)
        self.assertEqual(ctx.exception.code, 404)

    def test_invalid_request_id_returns_400(self) -> None:
        with self.assertRaises(urllib.error.HTTPError) as ctx:
            urllib.request.urlopen(self.url("/api/ai-requests/bad%20id/compliance"), timeout=30)
        self.assertEqual(ctx.exception.code, 400)

    def test_healthz(self) -> None:
        with urllib.request.urlopen(self.url("/healthz"), timeout=30) as response:
            self.assertEqual(json.load(response)["status"], "ok")


if __name__ == "__main__":
    unittest.main(verbosity=2)
