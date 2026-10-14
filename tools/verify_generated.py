"""Compare a generated checklist against the template.

Usage:
    python backend/tools/verify_generated.py <template.docx> <generated.docx>

Prints the answer matrix, every cell that differs from the template, and the
font used for the ballot-box runs, so the fill can be reviewed without a
renderer.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402

from compliance.catalog import (  # noqa: E402
    NA_COL,
    NO_COL,
    QUESTION_COL,
    YES_COL,
)
from compliance.generator import cell_text, iter_question_rows  # noqa: E402


def ballot_font(cell) -> str:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                name = run.font.name
                if name is None:
                    name = paragraph.style.font.name if paragraph.style else None
                return str(name)
    return "<none>"


def main(template: str, generated: str) -> int:
    template_doc = Document(template)
    generated_doc = Document(generated)

    template_rows = dict(iter_question_rows(template_doc))
    generated_rows = dict(iter_question_rows(generated_doc))

    print(f"template : {Path(template).name}")
    print(f"generated: {Path(generated).name}")
    print()
    print(f"{'id':<5}{'Yes':<6}{'No':<6}{'NA':<6}question")
    for question_id, row in generated_rows.items():
        yes = cell_text(row.cells[YES_COL])
        no = cell_text(row.cells[NO_COL])
        na = cell_text(row.cells[NA_COL])
        question = cell_text(row.cells[QUESTION_COL])
        print(f"{question_id:<5}{yes:<6}{no:<6}{na:<6}{question[:60]}")

    print()
    changes = []
    for question_id, row in generated_rows.items():
        original = template_rows.get(question_id)
        if original is None:
            changes.append(f"{question_id}: row missing from the template")
            continue
        for label, column in (("Yes", YES_COL), ("No", NO_COL), ("NA", NA_COL)):
            before = cell_text(original.cells[column])
            after = cell_text(row.cells[column])
            if before != after:
                changes.append(f"{question_id} {label}: {before!r} -> {after!r}")

    print(f"changed cells: {len(changes)}")
    for line in changes:
        print("  " + line)

    sample = next(iter(generated_rows.values()))
    print()
    print(f"font on the Yes/No/NA cells: {ballot_font(sample.cells[YES_COL])}")

    unanswered = [
        question_id
        for question_id, row in generated_rows.items()
        if cell_text(row.cells[YES_COL]).strip() not in {"\u2611", "\u2612"}
    ]
    if unanswered:
        print("NOT answered Yes: " + ", ".join(unanswered))
        return 1
    print("every question answered Yes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1], sys.argv[2]))
